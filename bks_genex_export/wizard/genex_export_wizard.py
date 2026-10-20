import base64
import io
import re
from datetime import datetime

from odoo import _, api, fields, models
from odoo.exceptions import UserError

try:
    import openpyxl
    from openpyxl.styles import Font
except ImportError as err:
    openpyxl = None
    _openpyxl_import_error = err
else:
    _openpyxl_import_error = None

_BKS_MOVE_GENEX_STATES = ("posted",)

GENEX_JVTICKET_HEADERS = (
    "COMP_CODE",
    "BRANCH_CODE",
    "CURRENCY_CODE",
    "GL_CODE",
    "CIF_SUB_NO",
    "SL_NO",
    "DEAL",
    "CLIENT",
    "VENDOR",
    "CV_AMOUNT",
    "FC_AMOUNT",
    "EXCH_RATE",
    "VALUE_DATE",
    "HEAD_DESC",
    "TR_CODE",
    "JV_TYPE",
    "TRANS_DATE",
    "DESCRIPTION",
    "ACC_REFERENCE",
    "processed",
    "COST_CENTER",
    "JV_REF",
    "LINK_TO_RJV_TICKET",
    "EMPTY",
    "VAT_CODE",
    "DIV_CODE",
    "DEPT_CODE",
    "Référence de lettrage",
)

IMAL_REGLEMENT_HEADERS = (
    "Code banque",
    "Montant",
    "Beneficiaire",
    "RIB",
    "Mode de règlement",
    "Référence règlement Odoo",
)


class BksGenexExportWizard(models.TransientModel):
    _name = "bks.genex.export.wizard"
    _description = "Export fichier GENEX"

    date_from = fields.Date(string="Du", required=True)
    date_to = fields.Date(string="Au", required=True)
    journal_ids = fields.Many2many("account.journal", string="Journaux")
    move_ids = fields.Many2many(
        comodel_name="account.move",
        relation="bks_genex_export_wizard_move_rel",
        column1="wizard_id",
        column2="move_id",
        string="Écritures à exporter",
        help="Retirez les lignes à exclure de l'export GENEX.",
    )
    payment_ids = fields.Many2many(
        comodel_name="account.payment",
        relation="bks_genex_export_wizard_payment_rel",
        column1="wizard_id",
        column2="payment_id",
        string="Paiements à exporter",
    )
    move_to_export_count = fields.Integer(
        string="Nombre d'écritures",
        compute="_compute_move_to_export_count",
    )
    payment_to_export_count = fields.Integer(
        string="Nombre de paiements",
        compute="_compute_move_to_export_count",
    )
    file_data = fields.Binary(string="Fichier", readonly=True)
    file_name = fields.Char(string="Nom du fichier", readonly=True)
    bks_payment_export = fields.Boolean(
        string="Export depuis paiements",
        default=lambda self: bool(self.env.context.get("bks_genex_payment_export")),
    )

    @api.model
    def default_get(self, fields_list):
        res = super().default_get(fields_list)
        active_model = self.env.context.get("active_model")
        active_ids = self.env.context.get("active_ids") or []
        if active_model == "account.payment" and active_ids:
            payments = self.env["account.payment"].browse(active_ids)
            payments._bks_check_genex_export_eligibility()
            res["payment_ids"] = [(6, 0, payments.ids)]
            res["bks_payment_export"] = True
            dates = payments.mapped("date")
            if dates:
                res["date_from"] = min(dates)
                res["date_to"] = max(dates)
        elif active_model == "account.move" and active_ids:
            moves = self.env["account.move"].browse(active_ids).filtered(
                lambda m: m.state in _BKS_MOVE_GENEX_STATES and not m.genex_exported
            )
            res["move_ids"] = [(6, 0, moves.ids)]
            dates = moves.mapped("date")
            if dates:
                res["date_from"] = min(dates)
                res["date_to"] = max(dates)
        if not res.get("date_from"):
            today = fields.Date.context_today(self)
            res["date_from"] = today.replace(day=1)
            res["date_to"] = today
        if (
            not res.get("bks_payment_export")
            and "move_ids" in fields_list
            and not res.get("move_ids")
        ):
            res["move_ids"] = [
                (6, 0, self._search_moves_for_criteria(res).ids)
            ]
        return res

    @api.model
    def _search_moves_for_criteria(self, vals=None):
        vals = vals or {}
        date_from = vals.get("date_from")
        date_to = vals.get("date_to")
        if not date_from or not date_to:
            return self.env["account.move"]
        journal_ids = vals.get("journal_ids")
        if journal_ids and journal_ids[0][0] == 6:
            journal_id_list = journal_ids[0][2]
        else:
            journal_id_list = []
        domain = [
            ("date", ">=", date_from),
            ("date", "<=", date_to),
            ("company_id", "in", self.env.companies.ids),
            ("state", "in", _BKS_MOVE_GENEX_STATES),
            ("genex_exported", "=", False),
        ]
        if journal_id_list:
            domain.append(("journal_id", "in", journal_id_list))
        return self.env["account.move"].search(domain, order="date, name, id")

    def _domain_moves_for_criteria(self):
        self.ensure_one()
        domain = [
            ("date", ">=", self.date_from),
            ("date", "<=", self.date_to),
            ("company_id", "in", self.env.companies.ids),
            ("state", "in", _BKS_MOVE_GENEX_STATES),
            ("genex_exported", "=", False),
        ]
        if self.journal_ids:
            domain.append(("journal_id", "in", self.journal_ids.ids))
        return domain

    def _is_payment_genex_export(self):
        self.ensure_one()
        return self.bks_payment_export

    def action_reload_moves(self):
        self.ensure_one()
        if self.date_from > self.date_to:
            raise UserError(_("La date de début doit être antérieure à la date de fin."))
        if self._is_payment_genex_export():
            raise UserError(
                _("Actualisez la sélection depuis la liste des paiements.")
            )
        moves = self.env["account.move"].search(
            self._domain_moves_for_criteria(), order="date, name, id"
        )
        self.move_ids = [(6, 0, moves.ids)]
        return self._reopen_wizard()

    @api.depends("move_ids", "payment_ids")
    def _compute_move_to_export_count(self):
        for wizard in self:
            wizard.move_to_export_count = len(wizard.move_ids)
            wizard.payment_to_export_count = len(wizard.payment_ids)

    def _reopen_wizard(self):
        return {
            "type": "ir.actions.act_window",
            "res_model": self._name,
            "view_mode": "form",
            "res_id": self.id,
            "target": "new",
            "name": _("Export GENEX"),
        }

    def _get_moves(self, raise_if_empty=True):
        self.ensure_one()
        moves = self.move_ids
        already = moves.filtered("genex_exported")
        if already:
            raise UserError(
                _(
                    "Certaines écritures sont déjà exportées GENEX : %(names)s",
                    names=", ".join(already.mapped("name")[:10]),
                )
            )
        invalid_state = moves.filtered(lambda m: m.state not in _BKS_MOVE_GENEX_STATES)
        if invalid_state:
            raise UserError(
                _(
                    "Seules les écritures comptabilisées peuvent être "
                    "exportées GENEX : %(names)s",
                    names=", ".join(invalid_state[:10].mapped("display_name")),
                )
            )
        moves = moves.filtered(lambda m: not m.genex_exported)
        if not moves and raise_if_empty:
            raise UserError(_("Aucune écriture à exporter."))
        return moves

    def _get_payments(self, raise_if_empty=True):
        self.ensure_one()
        payments = self.payment_ids
        payments._bks_check_genex_export_eligibility()
        already = payments.filtered("genex_exported")
        if already:
            raise UserError(
                _(
                    "Certains paiements sont déjà exportés GENEX : %(names)s",
                    names=", ".join(already.mapped("display_name")[:10]),
                )
            )
        if not payments and raise_if_empty:
            raise UserError(_("Aucun paiement à exporter."))
        return payments

    @staticmethod
    def _gl_code(account, company=None):
        """GL_CODE = code Odoo standard du compte (numérique si possible)."""
        if company:
            account = account.with_company(company)
        code = (account.code or "").strip()
        if code.isdigit():
            return int(code)
        return code

    def _account_has_genex_nomenclature(self, account, company):
        account = account.with_company(company)
        return bool(
            account.genex_agency_code
            and account.genex_currency_code
            and (account.code or "").strip()
        )

    def _genex_account_columns(self, line, company):
        """BRANCH_CODE, CURRENCY_CODE, GL_CODE, CIF_SUB_NO, SL_NO.

        Priorité : nomenclature du compte (agence, devise, code, CIF, séquence).
        Sinon repli analytique / société / partenaire (legacy).
        """
        account = line.account_id.with_company(company)
        if self._account_has_genex_nomenclature(account, company):
            return {
                "BRANCH_CODE": account.genex_agency_code,
                "CURRENCY_CODE": account.genex_currency_code,
                "GL_CODE": self._gl_code(account, company),
                "CIF_SUB_NO": account.genex_cif or 0,
                "SL_NO": account.genex_sequence or 0,
            }
        partner = line.move_id.partner_id
        cif = partner.genex_cif or 0 if partner else 0
        currency_code = company.genex_currency_code or 504
        return {
            "BRANCH_CODE": self._agency_code(line, company),
            "CURRENCY_CODE": currency_code,
            "GL_CODE": self._gl_code(account, company),
            "CIF_SUB_NO": cif,
            "SL_NO": 1 if cif else 0,
        }

    def _agency_code(self, line, company):
        account = line.account_id
        if account.genex_agency_code:
            return account.genex_agency_code
        distribution = line.analytic_distribution or {}
        if distribution:
            account_ids = [int(aid) for aid in distribution if str(aid).isdigit()]
            analytics = self.env["account.analytic.account"].browse(account_ids).exists()
            for analytic in analytics:
                if analytic.genex_agency_code:
                    return analytic.genex_agency_code
        return company.genex_default_agency_code or 1000

    def _payment_display_reference(self, payment):
        return (
            payment.name
            or payment.payment_reference
            or payment.memo
            or payment.display_name
        )

    def _short_reference(self, move, payment=None):
        fnp = move._bks_genex_fnp_record()
        if fnp:
            reception = move._bks_genex_reception_reference()
            base = fnp.name or (move.ref or move.name or "")
            if reception:
                return "%s %s" % (base, reception)
            return base
        if payment:
            pay_ref = self._payment_display_reference(payment)
            if pay_ref:
                return str(pay_ref).split()[0][:32]
        ref = (move.ref or move.payment_reference or move.name or "").strip()
        if not ref:
            return move.name or ""
        return ref.split()[0][:32]

    def _jv_ref(self, move):
        reception = move._bks_genex_reception_reference()
        if reception:
            return reception[:32]
        ref = (move.ref or move.payment_reference or move.name or "").strip()
        return ref.split()[0][:32] if ref else (move.name or "")[:32]

    def _long_label(self, move):
        fnp = move._bks_genex_fnp_record()
        if fnp and move._bks_genex_reception_reference():
            if fnp.reversal_move_id and fnp.reversal_move_id == move:
                return _("Extourne %s — réception %s") % (
                    fnp.name,
                    fnp.picking_id.name,
                )[:255]
            return _("FNP %s — réception %s") % (fnp.name, fnp.picking_id.name)[:255]
        ref = self._short_reference(move)
        partner = move.partner_id.name or ""
        invoice_no = move.name or move.ref or ""
        desc = (
            move.narration
            or move.invoice_origin
            or move.ref
            or (move.line_ids[:1].name if move.line_ids else "")
            or ""
        )
        if isinstance(desc, str):
            desc = re.sub(r"<[^>]+>", " ", desc)
        desc = " ".join(str(desc).split())
        parts = [p for p in (ref, partner, f"FN°{invoice_no}", desc) if p]
        return " ".join(parts)[:255]

    def _line_rows(self, move, company):
        rows = []
        jv_type = company.genex_ec_type or 500
        company_code = company.genex_company_code or 1
        head_desc = self._short_reference(move)
        jv_ref = self._jv_ref(move)
        lettrage_ref = move._bks_genex_reconciliation_reference()[:64]
        long_label = self._long_label(move)
        move_date = move.date
        rib = move._get_genex_rib()

        lines = move.line_ids.filtered(
            lambda line: line.display_type
            not in ("line_section", "line_subsection", "line_note")
            and line.account_id
        )
        for line in lines:
            account = line.account_id.with_company(company)
            amount_mad = line.balance
            if company.currency_id.is_zero(amount_mad):
                continue
            amount_currency = None
            rate = None
            if (
                line.currency_id
                and line.currency_id != company.currency_id
                and not line.currency_id.is_zero(line.amount_currency)
            ):
                amount_currency = line.amount_currency
                if line.currency_rate:
                    rate = line.currency_rate

            genex_keys = self._genex_account_columns(line, company)
            rows.append(
                {
                    "COMP_CODE": company_code,
                    **genex_keys,
                    "CV_AMOUNT": amount_mad,
                    "FC_AMOUNT": amount_currency,
                    "EXCH_RATE": rate,
                    "VALUE_DATE": move_date,
                    "HEAD_DESC": head_desc,
                    "JV_TYPE": jv_type,
                    "TRANS_DATE": move_date,
                    "DESCRIPTION": long_label,
                    "ACC_REFERENCE": rib,
                    "JV_REF": jv_ref,
                    "LINK_TO_RJV_TICKET": 0,
                    "Référence de lettrage": lettrage_ref or None,
                }
            )
        return rows

    @staticmethod
    def _genex_cell_value(value):
        if value is False or value is True:
            return None
        return value

    @staticmethod
    def _export_timestamp():
        return datetime.now().strftime("%Y%m%d_%H%M%S")

    def _store_filestore_attachment(self, filename, raw_bytes):
        return self.env["ir.attachment"].create(
            {
                "name": filename,
                "type": "binary",
                "raw": raw_bytes,
                "mimetype": (
                    "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet"
                ),
            }
        )

    def _build_jvticket_workbook(self, moves):
        if not openpyxl:
            raise UserError(
                _("Bibliothèque Python openpyxl requise : %s") % _openpyxl_import_error
            )
        wb = openpyxl.Workbook()
        ws = wb.active
        ws.title = "Sheet1"
        company = moves[:1].company_id if moves else self.env.company
        header_font = Font(bold=True)
        ws.cell(row=1, column=1, value=company.genex_company_code or 1)
        ws.cell(row=2, column=1, value=company.genex_default_agency_code or 1000)
        ws.cell(row=2, column=2, value="GENEX_JVTICKET")
        for col, title in enumerate(GENEX_JVTICKET_HEADERS, start=1):
            cell = ws.cell(row=4, column=col, value=title)
            cell.font = header_font

        row_idx = 5
        for move in moves:
            for row in self._line_rows(move, move.company_id):
                for col, key in enumerate(GENEX_JVTICKET_HEADERS, start=1):
                    value = row.get(key)
                    if isinstance(value, datetime):
                        value = value.date()
                    ws.cell(
                        row=row_idx,
                        column=col,
                        value=self._genex_cell_value(value),
                    )
                row_idx += 1
        # Marqueur de fin de fichier (modèle DAS_GENEX_JVTICKET : ligne vide puis <EOF>)
        ws.cell(row=row_idx + 1, column=1, value="<EOF>")
        return wb

    @staticmethod
    def _payment_bank_code(payment):
        journal = payment.journal_id
        bank = journal.bank_account_id
        if bank and bank.bank_id and bank.bank_id.bic:
            return bank.bank_id.bic
        if journal.code:
            return journal.code
        return ""

    def _payment_settlement_mode_label(self, payment):
        field = payment._fields.get("l10n_ma_reports_payment_method")
        code = payment.l10n_ma_reports_payment_method
        if field and code:
            labels = dict(field._description_selection(payment.env))
            label = labels.get(code)
            if label:
                return label
        if payment.payment_method_line_id:
            return payment.payment_method_line_id.name or ""
        return ""

    def _build_reglements_workbook(self, payments):
        if not openpyxl:
            raise UserError(
                _("Bibliothèque Python openpyxl requise : %s") % _openpyxl_import_error
            )
        wb = openpyxl.Workbook()
        ws = wb.active
        ws.title = "Reglements"
        header_font = Font(bold=True)
        for col, title in enumerate(IMAL_REGLEMENT_HEADERS, start=1):
            cell = ws.cell(row=1, column=col, value=title)
            cell.font = header_font
        row_idx = 2
        for payment in payments:
            rib = payment.partner_bank_id.sanitized_acc_number or payment.partner_bank_id.acc_number or ""
            rib = re.sub(r"[\s\-]", "", str(rib))[:24]
            ws.cell(row=row_idx, column=1, value=self._payment_bank_code(payment))
            ws.cell(row=row_idx, column=2, value=payment.amount)
            ws.cell(row=row_idx, column=3, value=payment.partner_id.name or "")
            ws.cell(row=row_idx, column=4, value=rib or None)
            ws.cell(
                row=row_idx,
                column=5,
                value=self._payment_settlement_mode_label(payment) or None,
            )
            ws.cell(
                row=row_idx,
                column=6,
                value=self._payment_display_reference(payment) or None,
            )
            row_idx += 1
        return wb

    def action_export_genex(self):
        self.ensure_one()
        if self.date_from > self.date_to:
            raise UserError(_("La date de début doit être antérieure à la date de fin."))

        ts = self._export_timestamp()
        if self._is_payment_genex_export():
            payments = self._get_payments()
            wb = self._build_reglements_workbook(payments)
            filename = "ODOO_IMAL_REGLEMENTS_%s.xlsx" % ts
            payments._mark_genex_exported()
        else:
            moves = self._get_moves()
            wb = self._build_jvticket_workbook(moves)
            filename = "ODOO_GENEX_%s.xlsx" % ts
            moves._mark_genex_exported()

        buffer = io.BytesIO()
        wb.save(buffer)
        raw = buffer.getvalue()
        self._store_filestore_attachment(filename, raw)
        self.write(
            {
                "file_data": base64.b64encode(raw),
                "file_name": filename,
            }
        )
        return self._reopen_wizard()
