#!/usr/bin/env python3
"""Export taxes BKS (TVA + retenue à la source) vers Excel import Odoo 19."""
from __future__ import annotations

import json
import re
import subprocess
import textwrap
from pathlib import Path

import openpyxl
from openpyxl.styles import Font
from openpyxl.utils import get_column_letter

JSON_CACHE = Path("/tmp/bks_taxes.json")
OUTPUT_PATHS = [
    Path("/Users/macbookpro/Documents/osisoftware/bks/assets/Import_taxes_Odoo19_BKS.xlsx"),
    Path("/Users/macbookpro/odoo19/bks/assets/Import_taxes_Odoo19_BKS.xlsx"),
]

IMPORT_HEADERS = [
    "id",
    "name",
    "description",
    "amount",
    "amount_type",
    "type_tax_use",
    "tax_group_id",
    "sequence",
    "code",
    "active",
    "tax_exigibility",
    "is_ras",
    "is_ras_tva",
    "cash_basis_transition_account_id",
    "repartition_line_ids/repartition_type",
    "repartition_line_ids/document_type",
    "repartition_line_ids/account_id",
    "repartition_line_ids/factor_percent",
]

REPART_HEADERS = [
    "tax_name",
    "document_type",
    "repartition_type",
    "factor_percent",
    "account_code",
    "account_name",
]


def slugify(name: str) -> str:
    s = re.sub(r"[^a-zA-Z0-9]+", "_", name.lower()).strip("_")
    return s[:60] or "tax"


def fetch_taxes() -> list[dict]:
    if JSON_CACHE.exists():
        return json.loads(JSON_CACHE.read_text(encoding="utf-8"))

    script = textwrap.dedent(
        '''
        import json, re
        Tax = env["account.tax"].with_context(lang="fr_FR", active_test=False)
        company = env.company
        out = []
        for tax in Tax.search([("company_id", "=", company.id)], order="type_tax_use, sequence, id"):
            def repr_lines(lines):
                res = []
                for l in lines.sorted(key=lambda x: (x.document_type, x.sequence, x.id)):
                    res.append({
                        "document_type": l.document_type,
                        "repartition_type": l.repartition_type,
                        "factor_percent": l.factor_percent,
                        "account_code": l.account_id.code if l.account_id else "",
                        "account_name": l.account_id.name if l.account_id else "",
                    })
                return res
            desc = tax.description
            if desc:
                desc = re.sub(r"<[^>]+>", "", str(desc))
            out.append({
                "id": tax.id,
                "name": tax.name,
                "description": desc or "",
                "amount": tax.amount,
                "amount_type": tax.amount_type,
                "type_tax_use": tax.type_tax_use,
                "sequence": tax.sequence,
                "code": tax.code or "",
                "tax_group": tax.tax_group_id.name or "",
                "active": tax.active,
                "tax_exigibility": tax.tax_exigibility or "",
                "is_ras": bool(tax.is_ras),
                "is_ras_tva": bool(tax.is_ras_tva),
                "cash_basis_account": tax.cash_basis_transition_account_id.code if tax.cash_basis_transition_account_id else "",
                "repartition_lines": repr_lines(tax.repartition_line_ids),
            })
        print(json.dumps(out, ensure_ascii=False))
        '''
    )
    proc = subprocess.run(
        [
            "/Users/macbookpro/.pyenv/versions/odoo18/bin/python3",
            "odoo-bin",
            "shell",
            "-d",
            "bks",
            "-c",
            "odoo.conf",
            "--addons-path=/Users/macbookpro/odoo19/odoo/odoo/addons,/Users/macbookpro/odoo19/odoo/addons,/Users/macbookpro/odoo19/odoo/enterprise,/Users/macbookpro/odoo19/osicompta,/Users/macbookpro/odoo19/bks",
            "--no-http",
        ],
        input=script,
        cwd="/Users/macbookpro/odoo19/odoo",
        capture_output=True,
        text=True,
    )
    for line in proc.stdout.splitlines():
        line = line.strip()
        if line.startswith("["):
            data = json.loads(line)
            JSON_CACHE.write_text(json.dumps(data, ensure_ascii=False, indent=2), encoding="utf-8")
            return data
    raise RuntimeError(proc.stderr[-4000:])


def tax_base_row(tax: dict) -> list:
    ext_id = f"bks_tax_{slugify(tax['name'])}"
    return [
        ext_id,
        tax["name"],
        tax["description"],
        tax["amount"],
        tax["amount_type"],
        tax["type_tax_use"],
        tax["tax_group"],
        tax["sequence"],
        tax["code"],
        tax["active"],
        tax["tax_exigibility"],
        tax["is_ras"],
        tax["is_ras_tva"],
        tax["cash_basis_account"],
        "",
        "",
        "",
        "",
    ]


def repart_row(tax: dict, line: dict) -> list:
    return [
        "",
        "",
        "",
        "",
        "",
        "",
        "",
        "",
        "",
        "",
        "",
        "",
        "",
        "",
        line["repartition_type"],
        line["document_type"],
        line["account_code"],
        line["factor_percent"],
    ]


def write_workbook(taxes: list[dict], path: Path) -> None:
    path.parent.mkdir(parents=True, exist_ok=True)
    wb = openpyxl.Workbook()

    # --- Synthese ---
    ws = wb.active
    ws.title = "Synthese"
    ws["B2"] = "Export taxes Odoo 19 — base BKS"
    ws["B2"].font = Font(bold=True, size=14)
    rows = [
        ("Source", "Base PostgreSQL bks + module account_tax_code"),
        ("Taxes exportées", len(taxes)),
        ("TVA actives", sum(1 for t in taxes if t["active"] and not t["is_ras"] and not t["is_ras_tva"])),
        ("Retenue à la source (is_ras)", sum(1 for t in taxes if t["is_ras"])),
        ("RAS TVA (is_ras_tva)", sum(1 for t in taxes if t["is_ras_tva"])),
        ("Onglet import", "Import_Odoo19"),
    ]
    for i, (k, v) in enumerate(rows, start=4):
        ws.cell(row=i, column=2, value=k)
        ws.cell(row=i, column=3, value=v)

    # --- Import Odoo 19 (format multi-lignes repartition) ---
    ws_imp = wb.create_sheet("Import_Odoo19")
    ws_imp.append(IMPORT_HEADERS)
    for tax in taxes:
        lines = tax["repartition_lines"] or [
            {
                "document_type": "invoice",
                "repartition_type": "base",
                "factor_percent": 100.0,
                "account_code": "",
            }
        ]
        first = True
        for line in lines:
            if first:
                row = tax_base_row(tax)
                row[14] = line["repartition_type"]
                row[15] = line["document_type"]
                row[16] = line["account_code"]
                row[17] = line["factor_percent"]
                ws_imp.append(row)
                first = False
            else:
                ws_imp.append(repart_row(tax, line))

    # --- Retenue à la source ---
    ws_ras = wb.create_sheet("Retenue_Source")
    ras_headers = [
        "name",
        "amount",
        "type_tax_use",
        "code",
        "is_ras",
        "is_ras_tva",
        "tax_group_id",
        "compte_taxe_facture",
        "compte_taxe_avoir",
    ]
    ws_ras.append(ras_headers)
    for tax in taxes:
        if not (tax["is_ras"] or tax["is_ras_tva"]):
            continue
        inv_acc = next(
            (
                l["account_code"]
                for l in tax["repartition_lines"]
                if l["document_type"] == "invoice" and l["repartition_type"] == "tax"
            ),
            "",
        )
        ref_acc = next(
            (
                l["account_code"]
                for l in tax["repartition_lines"]
                if l["document_type"] == "refund" and l["repartition_type"] == "tax"
            ),
            "",
        )
        ws_ras.append(
            [
                tax["name"],
                tax["amount"],
                tax["type_tax_use"],
                tax["code"],
                tax["is_ras"],
                tax["is_ras_tva"],
                tax["tax_group"],
                inv_acc,
                ref_acc,
            ]
        )

    # --- Repartition detail ---
    ws_rep = wb.create_sheet("Repartition_Detail")
    ws_rep.append(REPART_HEADERS)
    for tax in taxes:
        for line in tax["repartition_lines"]:
            ws_rep.append(
                [
                    tax["name"],
                    line["document_type"],
                    line["repartition_type"],
                    line["factor_percent"],
                    line["account_code"],
                    line["account_name"],
                ]
            )

    # --- Instructions ---
    ws_inst = wb.create_sheet("Instructions")
    instructions = [
        "Import des taxes dans Odoo 19 (BKS / Maroc)",
        "",
        "1. Onglet à importer : « Import_Odoo19 » (format multi-lignes avec répartition).",
        "2. Champs RÀS : is_ras et is_ras_tva (module account_tax_code requis).",
        "3. Comptes : repartition_line_ids/account_id = code comptable à 6 chiffres du plan GL.",
        "4. Retenue à la source : voir l’onglet « Retenue_Source » pour contrôle rapide.",
        "5. Import Odoo : Comptabilité → Configuration → Taxes → Favoris → Importer un fichier.",
        "6. Mapper les colonnes techniques ; tester puis valider l’import.",
        "7. Les taxes TVA marocaines utilisent tax_exigibility=on_payment et compte 445510 (encaissement).",
    ]
    for i, line in enumerate(instructions, start=2):
        ws_inst.cell(row=i, column=2, value=line)

    for sheet in wb.worksheets:
        for col in range(1, sheet.max_column + 1):
            letter = get_column_letter(col)
            max_len = 0
            for cell in sheet[letter]:
                if cell.value is not None:
                    max_len = max(max_len, len(str(cell.value)))
            sheet.column_dimensions[letter].width = min(max(max_len + 2, 10), 48)

    wb.save(path)


def main():
    taxes = fetch_taxes()
    written = []
    for path in OUTPUT_PATHS:
        write_workbook(taxes, path)
        written.append(str(path))
    print({"taxes": len(taxes), "files": written})


if __name__ == "__main__":
    main()
