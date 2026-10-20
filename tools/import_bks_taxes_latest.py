#!/usr/bin/env python3
"""Import taxes BKS + mise à jour taxes achat articles (latest assets)."""
from __future__ import annotations

import importlib.util
import re
import sys
from pathlib import Path

import openpyxl

TAXES_PATH = Path(
    "/Users/macbookpro/Documents/osisoftware/bks/assets/latest/Taxes_Odoo19_BKS.xlsx"
)
PRODUCT_PATH = Path(
    "/Users/macbookpro/Documents/osisoftware/bks/assets/latest/Référentiel_Articles_ODOO.xlsx"
)

FIRST_TAX_ROW = 3


def _load_ref_helpers():
    path = Path(__file__).parent / "import_bks_latest_referentiels.py"
    spec = importlib.util.spec_from_file_location("ref", path)
    mod = importlib.util.module_from_spec(spec)
    spec.loader.exec_module(mod)
    return mod


def _parse_amount(raw):
    if raw is None:
        return 0.0
    if isinstance(raw, (int, float)):
        return float(raw)
    s = str(raw).strip()
    m = re.search(r"[-+]?\d+(?:\.\d+)?", s)
    return float(m.group()) if m else 0.0


def _as_bool(val):
    if isinstance(val, bool):
        return val
    if val is None:
        return False
    return str(val).strip().lower() in ("true", "1", "oui", "yes")


def _normalize_tax_name(name):
    if not name:
        return ""
    name = str(name).strip()
    name = re.sub(r"\)+$", "", name)
    return name.strip()


def _split_ras_names(cell):
    if not cell:
        return []
    parts = re.split(r",\s*", str(cell))
    return [_normalize_tax_name(p) for p in parts if _normalize_tax_name(p)]


def _rmg_withholding(name, amount, is_ras, is_ras_tva):
    vals = {"is_withholding_tax_on_payment": False, "rmg_withholding_type": False}
    if is_ras_tva:
        vals["is_withholding_tax_on_payment"] = True
        if abs(amount) >= 19.9 or amount == -20:
            vals["rmg_withholding_type"] = "vat_100"
        elif amount in (-15, -7.5):
            vals["rmg_withholding_type"] = "vat_75"
        else:
            vals["rmg_withholding_type"] = "vat_100"
        return vals
    if is_ras:
        vals["is_withholding_tax_on_payment"] = True
        low = name.lower()
        if "5%" in name and "loyer" in low:
            vals["rmg_withholding_type"] = "cit_5"
        elif amount == -5 or ("5%" in name and "rémun" in low):
            vals["rmg_withholding_type"] = "cit_5"
        elif "30%" in name:
            vals["rmg_withholding_type"] = "pit_30"
        elif "10%" in name:
            vals["rmg_withholding_type"] = "pit_10"
        elif amount == -15:
            vals["rmg_withholding_type"] = "pit_10"
        return vals
    return vals


def _account_from_row(env, company, row, start, ref_mod):
    key = ref_mod._nomenclature_key(
        row[start], row[start + 1], row[start + 2], row[start + 3], row[start + 4]
    )
    fc = ref_mod._full_code(key)
    acc = env["account.account"].search(
        [
            ("genex_agency_code", "=", key[0]),
            ("genex_currency_code", "=", key[1]),
            ("code", "=", str(key[2])),
            ("genex_cif", "=", key[3]),
            ("genex_sequence", "=", key[4]),
            ("company_ids", "in", company.id),
        ],
        limit=1,
    )
    if not acc:
        raise ValueError(f"Compte taxe introuvable : {fc}")
    return acc


def _repartition_commands(account, reverse_charge=False):
    lines = [(0, 0, {"repartition_type": "base", "factor_percent": 100})]
    if reverse_charge:
        lines.append(
            (0, 0, {"repartition_type": "tax", "factor_percent": 100, "account_id": account.id})
        )
        lines.append(
            (0, 0, {"repartition_type": "tax", "factor_percent": -100, "account_id": account.id})
        )
    else:
        lines.append(
            (0, 0, {"repartition_type": "tax", "factor_percent": 100, "account_id": account.id})
        )
    return lines


def load_tax_rows(path=TAXES_PATH):
    wb = openpyxl.load_workbook(path, read_only=True, data_only=True)
    ws = wb["Retenue_Source & TVA"]
    rows = []
    for i, row in enumerate(ws.iter_rows(values_only=True), start=1):
        if i < FIRST_TAX_ROW or not row or not row[0]:
            continue
        rows.append(row)
    wb.close()
    return rows


def import_taxes(env, company, commit=False):
    ref_mod = _load_ref_helpers()
    Tax = env["account.tax"].with_context(
        tracking_disable=True, mail_create_nolog=True, mail_notrack=True
    )
    Group = env["account.tax.group"]
    by_name = {}

    created = updated = 0
    errors = []

    for row in load_tax_rows():
        name = str(row[0]).strip()
        try:
            amount = _parse_amount(row[1])
            is_ras = _as_bool(row[5])
            is_ras_tva = _as_bool(row[6])
            group_name = (row[7] or "Taxes").strip()
            reverse = "reverser" in name.lower() or "neutre" in str(row[1]).lower()

            inv_acc = _account_from_row(env, company, row, 9, ref_mod)
            ref_acc = _account_from_row(env, company, row, 15, ref_mod)

            group = Group.search(
                [("name", "=", group_name), ("company_id", "=", company.id)], limit=1
            )
            if not group:
                group = Group.create({"name": group_name, "company_id": company.id})

            tax_amount = abs(amount) if reverse else amount
            vals = {
                "name": name,
                "amount_type": "percent",
                "amount": tax_amount,
                "type_tax_use": "purchase",
                "company_id": company.id,
                "tax_group_id": group.id,
                "description": (row[3] or "") if row[3] else False,
                "active": True,
                **_rmg_withholding(name, amount, is_ras, is_ras_tva),
                "invoice_repartition_line_ids": [(5, 0, 0)]
                + _repartition_commands(inv_acc, reverse_charge=reverse),
                "refund_repartition_line_ids": [(5, 0, 0)]
                + _repartition_commands(ref_acc, reverse_charge=reverse),
            }

            tax = Tax.search(
                [("name", "=", name), ("company_id", "=", company.id)], limit=1
            )
            if tax:
                tax.write(vals)
                updated += 1
            else:
                tax = Tax.create(vals)
                created += 1
            by_name[name] = tax
            by_name[_normalize_tax_name(name)] = tax
        except Exception as exc:  # noqa: BLE001
            errors.append(f"{name}: {exc}")

    if commit:
        env.cr.commit()
        env.invalidate_all()
    return {
        "taxes_created": created,
        "taxes_updated": updated,
        "tax_errors": errors,
        "tax_by_name": by_name,
    }


def _find_tax(by_name, Tax, company, label):
    label = _normalize_tax_name(label)
    if not label or label in ("0%", "0"):
        return Tax
    if label in by_name:
        return by_name[label]
    tax = Tax.search(
        [
            ("name", "=", label),
            ("company_id", "=", company.id),
            ("type_tax_use", "=", "purchase"),
        ],
        limit=1,
    )
    if tax:
        return tax
    tax = Tax.search(
        [
            ("name", "ilike", label.replace(")", "").replace("(", "")[:40]),
            ("company_id", "=", company.id),
            ("type_tax_use", "=", "purchase"),
        ],
        limit=1,
    )
    return tax


def update_product_taxes(env, company, by_name, commit=False):
    ref_mod = _load_ref_helpers()
    wb = openpyxl.load_workbook(PRODUCT_PATH, read_only=True, data_only=True)
    ws = wb.active
    Tax = env["account.tax"]
    Product = env["product.template"].with_context(
        tracking_disable=True, mail_create_nolog=True, mail_notrack=True
    )
    updated = skipped = 0
    errors = []

    for row in ws.iter_rows(min_row=2, values_only=True):
        ref = ref_mod._clean(row[0])
        if not ref:
            continue
        try:
            tmpl = Product.search([("default_code", "=", ref)], limit=1)
            if not tmpl:
                skipped += 1
                continue
            tax_ids = []
            achat = _normalize_tax_name(row[3])
            if achat and achat not in ("0%", "0"):
                t = _find_tax(by_name, Tax, company, achat)
                if not t:
                    raise ValueError(f"Taxe achat introuvable : {achat}")
                tax_ids.extend(
                    t.filtered(lambda tax: not tax.is_withholding_tax_on_payment).ids
                )
            # Les taxes RAS ne sont pas sur l'article : le moteur RMG les propose sur la facture.
            tmpl.write({"supplier_taxes_id": [(6, 0, list(dict.fromkeys(tax_ids)))]})
            updated += 1
        except Exception as exc:  # noqa: BLE001
            errors.append(f"{ref}: {exc}")
    wb.close()
    if commit:
        env.cr.commit()
    return {
        "product_taxes_updated": updated,
        "product_taxes_skipped": skipped,
        "product_tax_errors": errors[:30],
        "product_tax_error_count": len(errors),
    }


def strip_product_withholding_taxes(env, company_id=None, commit=False):
    """Retire les taxes RAS des articles (TVA achat uniquement)."""
    company = env["res.company"].browse(company_id or 1)
    Product = env["product.template"].with_context(
        tracking_disable=True, mail_create_nolog=True, mail_notrack=True
    )
    updated = 0
    for tmpl in Product.search([("purchase_ok", "=", True)]):
        purchase_taxes = tmpl.supplier_taxes_id.filtered(
            lambda tax: tax.company_id.id in (False, company.id)
        )
        regular = purchase_taxes.filtered(
            lambda tax: not tax.is_withholding_tax_on_payment
        )
        if purchase_taxes != regular:
            tmpl.write({"supplier_taxes_id": [(6, 0, regular.ids)]})
            updated += 1
    if commit:
        env.cr.commit()
    return {"products_withholding_stripped": updated}


def run_all(env, company_id=None, commit=False):
    company = env["res.company"].browse(company_id or 1)
    env = env(context=dict(env.context, allowed_company_ids=company.ids))
    tax_res = import_taxes(env, company, commit=commit)
    by_name = tax_res.pop("tax_by_name")
    prod_res = update_product_taxes(env, company, by_name, commit=commit)
    strip_res = strip_product_withholding_taxes(env, company.id, commit=commit)
    return {"company": company.name, **tax_res, **prod_res, **strip_res}


if __name__ == "__main__":
    print("Exécuter via odoo-bin shell.", file=sys.stderr)
    sys.exit(1)
