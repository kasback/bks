#!/usr/bin/env python3
"""Import referentiels BKS (latest) : plan comptable, fournisseurs, articles.

Fichiers attendus dans assets/latest/ :
  - Plan_Comptable_ODOO.xlsx
  - Référentiel_Fournisseur_ODOO.xlsx (onglet Feuil1)
  - Référentiel_Articles_ODOO.xlsx

Usage (odoo-bin shell -d bks_prod --no-http) :

    import importlib.util
    spec = importlib.util.spec_from_file_location(
        "imp", "/Users/macbookpro/odoo19/bks/tools/import_bks_latest_referentiels.py")
    m = importlib.util.module_from_spec(spec); spec.loader.exec_module(m)
    print(m.run_all(env, commit=True))
"""
from __future__ import annotations

import importlib.util
import re
import sys
from pathlib import Path

import openpyxl

ASSETS = Path("/Users/macbookpro/Documents/osisoftware/bks/assets/latest")
PLAN_PATH = ASSETS / "Plan_Comptable_ODOO.xlsx"
SUPPLIER_PATH = ASSETS / "Référentiel_Fournisseur_ODOO.xlsx"
PRODUCT_PATH = ASSETS / "Référentiel_Articles_ODOO.xlsx"

DEFAULT_AGENCY = 1000
DEFAULT_CURRENCY = 504
BATCH = 50

COUNTRY_CODES = {
    "MAROC": "MA",
    "EMIRATS ARABES UNIS": "AE",
    "BELGIQUE": "BE",
    "SUISSE": "CH",
    "FRANCE": "FR",
}

ACCOUNT_BLOCKS = (
    ("property_account_payable_id", 16),
    ("property_account_advance_id", 21),
    ("property_account_fnp_id", 26),
)


def _clean(value):
    if value is None:
        return None
    value = str(value).strip()
    return value or None


def _parse_agency(value):
    if value is None:
        return DEFAULT_AGENCY
    if isinstance(value, (int, float)):
        return int(value)
    s = str(value).strip()
    if s.isdigit():
        return int(s)
    return DEFAULT_AGENCY


def _nomenclature_key(agency, currency, gl, cif, seq):
    return (
        _parse_agency(agency),
        int(currency),
        str(int(gl)),
        int(cif or 0),
        int(seq or 0),
    )


def _full_code(key):
    a, c, gl, cif, seq = key
    return ".".join(str(x) for x in (a, c, gl, cif, seq))


def _account_type_and_reconcile(gl: int) -> tuple[str, bool]:
    if gl in (367104, 367106):
        return "liability_payable", True
    if gl == 359910:
        return "asset_current", False
    if str(gl).startswith("35") or str(gl).startswith("36"):
        return "asset_current", False
    if gl in (451401, 453101) or 451000 <= gl < 454000:
        return "asset_non_current", False
    if 461000 <= gl < 466000:
        return "asset_fixed", False
    if 643000 <= gl < 668000:
        return "expense", False
    return "asset_current", False


def load_plan_rows(path: Path = PLAN_PATH):
    wb = openpyxl.load_workbook(path, read_only=True, data_only=True)
    ws = wb.active
    rows = []
    for i, row in enumerate(ws.iter_rows(values_only=True), start=1):
        if i == 1:
            continue
        if not row or row[2] is None:
            continue
        rows.append(row)
    wb.close()
    return rows


def import_plan_accounts(env, company, commit=False):
    Account = env["account.account"].with_company(company).with_context(
        tracking_disable=True, mail_create_nolog=True, mail_notrack=True
    )
    ctx = dict(env.context, defer_account_code_checks=True)
    index: dict[tuple, object] = {}
    by_full: dict[str, object] = {}

    for acc in Account.with_context(active_test=False).search(
        [("company_ids", "in", company.id)]
    ):
        if acc.genex_full_code:
            by_full[acc.genex_full_code] = acc

    created = updated = 0
    errors = []
    for row in load_plan_rows():
        try:
            key = _nomenclature_key(row[0], row[1], row[2], row[3], row[4])
            fc = _full_code(key)
            name = (row[5] or fc).strip()
            gl = int(key[2])
            acc_type, reconcile = _account_type_and_reconcile(gl)
            acc = by_full.get(fc)
            vals = {
                "name": name,
                "genex_agency_code": key[0],
                "genex_currency_code": key[1],
                "genex_cif": key[3],
                "genex_sequence": key[4],
            }
            if acc:
                acc.with_context(**ctx).write(vals)
                updated += 1
            else:
                siblings = Account.search(
                    [
                        ("code", "=", str(gl)),
                        ("genex_agency_code", "=", key[0]),
                        ("genex_currency_code", "=", key[1]),
                        ("genex_cif", "=", key[3]),
                        ("genex_sequence", "=", key[4]),
                        ("company_ids", "in", company.id),
                    ],
                    limit=1,
                )
                if siblings:
                    acc = siblings
                    acc.with_context(**ctx).write(vals)
                    updated += 1
                else:
                    acc = Account.with_context(**ctx).create(
                        {
                            **vals,
                            "code": str(gl),
                            "account_type": acc_type,
                            "reconcile": reconcile,
                            "company_ids": [(6, 0, [company.id])],
                        }
                    )
                    created += 1
            index[key] = acc
            by_full[fc] = acc
        except Exception as exc:  # noqa: BLE001
            errors.append(f"{row}: {exc}")
        if commit and (created + updated) % BATCH == 0:
            env.cr.commit()
            env.invalidate_all()
    if commit:
        env.cr.commit()
        env.invalidate_all()
    return {"plan_created": created, "plan_updated": updated, "plan_errors": errors, "index": index}


def _load_supplier_module():
    path = Path(__file__).parent / "import_referentiel_fournisseur_bks.py"
    spec = importlib.util.spec_from_file_location("ref_f", path)
    mod = importlib.util.module_from_spec(spec)
    spec.loader.exec_module(mod)
    return mod


def import_suppliers(env, company, account_index, commit=False):
    ref = _load_supplier_module()
    ref.XLSX_PATH = SUPPLIER_PATH

    wb = openpyxl.load_workbook(SUPPLIER_PATH, read_only=True, data_only=True)
    ws = wb["Feuil1"]
    rows = []
    for i, row in enumerate(ws.iter_rows(values_only=True), start=1):
        if i < 3 or not row or not _clean(row[0]):
            continue
        rows.append(row)
    wb.close()

    Partner = env["res.partner"].with_context(
        tracking_disable=True, mail_create_nolog=True, mail_notrack=True
    )
    countries = {
        c.code: c
        for c in env["res.country"].search([("code", "in", list(COUNTRY_CODES.values()))])
    }
    currencies = {
        c.name: c for c in env["res.currency"].with_context(active_test=False).search([])
    }

    def full_code_row(row, start):
        return _full_code(
            _nomenclature_key(
                row[start], row[start + 1], row[start + 2], row[start + 3], row[start + 4]
            )
        )

    def account_for_row(row, start):
        key = _nomenclature_key(
            row[start], row[start + 1], row[start + 2], row[start + 3], row[start + 4]
        )
        acc = account_index.get(key)
        if not acc:
            fc = _full_code(key)
            acc = env["account.account"].search(
                [("genex_full_code", "=", fc), ("company_ids", "in", company.id)], limit=1
            )
        if not acc:
            raise ValueError(f"Compte manquant pour {_full_code(key)}")
        return acc

    created_partners = updated_partners = 0
    errors = []
    invalid_vat = []
    for n, row in enumerate(rows, start=1):
        ref_code = _clean(row[1])
        try:
            with env.cr.savepoint():
                vals, currency = ref._partner_vals(env, row, countries, currencies)
                partner = Partner.search([("ref", "=", ref_code)], limit=1) if ref_code else Partner
                if not partner:
                    partner = Partner.search(
                        [("name", "=ilike", vals["name"]), ("supplier_rank", ">", 0)], limit=1
                    )
                if vals.get("vat") and not re.fullmatch(
                    r"\d{7,8}", (vals["vat"] or "").replace(" ", "")
                ):
                    invalid_vat.append(f"{ref_code}: IF {vals['vat']}")
                    vals["comment"] = f"IF (référentiel, non conforme) : {vals['vat']}"
                    vals["vat"] = False
                if partner:
                    partner.write(vals)
                    updated_partners += 1
                else:
                    partner = Partner.create(vals)
                    created_partners += 1
                prop_vals = {
                    field: account_for_row(row, start).id
                    for field, start in ACCOUNT_BLOCKS
                }
                if currency:
                    prop_vals["property_purchase_currency_id"] = currency.id
                partner.with_company(company).write(prop_vals)
        except Exception as exc:  # noqa: BLE001
            errors.append(f"{ref_code} {row[0]}: {exc}")
        if commit and n % BATCH == 0:
            env.cr.commit()
    if commit:
        env.cr.commit()
    return {
        "suppliers_created": created_partners,
        "suppliers_updated": updated_partners,
        "invalid_vat": invalid_vat,
        "supplier_errors": errors,
    }


def _product_type(row):
    label = (_clean(row[7]) or "").lower()
    if "service" in label:
        return "service"
    if _clean(row[8]) and str(row[8]).lower() in ("true", "1", "oui", "yes"):
        return "consu"
    return "consu"


def import_products(env, company, account_index, commit=False):
    wb = openpyxl.load_workbook(PRODUCT_PATH, read_only=True, data_only=True)
    ws = wb.active
    Product = env["product.template"].with_context(
        tracking_disable=True, mail_create_nolog=True, mail_notrack=True
    )
    Tax = env["account.tax"].with_company(company)
    Category = env["product.category"]

    categ_cache: dict[str, object] = {}
    created = updated = skipped = 0
    errors = []

    def get_categ(name, parent_name=None):
        name = (name or "Divers").strip()
        if name in categ_cache:
            return categ_cache[name]
        parent = get_categ(parent_name) if parent_name else False
        cat = Category.search([("name", "=", name), ("parent_id", "=", parent.id if parent else False)], limit=1)
        if not cat:
            cat = Category.create({"name": name, "parent_id": parent.id if parent else False})
        categ_cache[name] = cat
        return cat

    def find_tax(label):
        if not label:
            return False
        tax = Tax.search(
            [("name", "=", label.strip()), ("type_tax_use", "=", "purchase")], limit=1
        )
        if not tax:
            tax = Tax.search([("name", "ilike", label.strip()[:20])], limit=1)
        return tax[:1]

    def expense_account(row):
        key = _nomenclature_key(
            row[13], row[14], row[15], row[16], row[17]
        )
        acc = account_index.get(key)
        if not acc:
            fc = _full_code(key)
            acc = env["account.account"].search(
                [("genex_full_code", "=", fc), ("company_ids", "in", company.id)], limit=1
            )
        return acc

    for n, row in enumerate(ws.iter_rows(min_row=2, values_only=True), start=1):
        ref = _clean(row[0])
        if not ref:
            continue
        try:
            with env.cr.savepoint():
                name = _clean(row[1]) or ref
                categ = get_categ(_clean(row[4]) or _clean(row[11]), _clean(row[10]))
                acc = expense_account(row)
                if not acc:
                    raise ValueError(f"Compte charge introuvable pour {ref}")
                tax = find_tax(row[3])
                vals = {
                    "name": name,
                    "default_code": ref,
                    "type": _product_type(row),
                    "categ_id": categ.id,
                    "purchase_ok": True,
                    "sale_ok": False,
                }
                if tax:
                    vals["supplier_taxes_id"] = [(6, 0, tax.ids)]
                tmpl = Product.search([("default_code", "=", ref)], limit=1)
                if tmpl:
                    tmpl.write(vals)
                    updated += 1
                else:
                    tmpl = Product.create(vals)
                    created += 1
                tmpl.with_company(company).write(
                    {"property_account_expense_id": acc.id}
                )
        except Exception as exc:  # noqa: BLE001
            errors.append(f"{ref}: {exc}")
        if commit and n % BATCH == 0:
            env.cr.commit()
    wb.close()
    if commit:
        env.cr.commit()
    return {
        "products_created": created,
        "products_updated": updated,
        "product_errors": errors[:20],
        "product_error_count": len(errors),
    }


def run_all(env, company_id=None, commit=False):
    company = env["res.company"].browse(company_id or 1)
    env = env(context=dict(env.context, allowed_company_ids=company.ids))
    plan = import_plan_accounts(env, company, commit=commit)
    index = plan.pop("index")
    suppliers = import_suppliers(env, company, index, commit=commit)
    products = import_products(env, company, index, commit=commit)
    return {"company": company.name, **plan, **suppliers, **products}


if __name__ == "__main__":
    print("Exécuter via odoo-bin shell.", file=sys.stderr)
    sys.exit(1)
