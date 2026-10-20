#!/usr/bin/env python3
"""Import du référentiel fournisseurs BKS (Référentiel Fournisseur ODOO.xlsx).

Pour chaque fournisseur :
  * création / mise à jour de la fiche (res.partner, supplier_rank = 1) ;
  * création (ou réutilisation) des 3 comptes comptables de la nomenclature
    ``agence.devise.code GL.CIF.séquence`` :
        - Compte fournisseur
        - Compte Avance fournisseur
        - Compte FNP fournisseur
  * affectation de ces comptes sur la fiche fournisseur.

Le script est idempotent (clé fournisseur = Référence ; clé compte = code GL +
agence + devise + CIF + séquence).

Exécution : dans ``odoo-bin shell -d <base> --no-http``, charger ce module puis
appeler ``run(env, commit=True)``.
"""
from __future__ import annotations

import logging
import re
import sys
from pathlib import Path

import openpyxl

XLSX_PATH = Path(
    "/Users/macbookpro/Documents/osisoftware/bks/assets/Référentiel Fournisseur ODOO.xlsx"
)
FIRST_DATA_ROW = 3  # ligne 1 : groupes, ligne 2 : en-têtes
BATCH_SIZE = 50

COUNTRY_CODES = {
    "MAROC": "MA",
    "EMIRATS ARABES UNIS": "AE",
    "BELGIQUE": "BE",
    "SUISSE": "CH",
    "FRANCE": "FR",
}

# (nom du champ partenaire, colonne début bloc compte)
ACCOUNT_BLOCKS = (
    ("property_account_payable_id", 16, "payable"),
    ("property_account_advance_id", 21, "advance"),
    ("property_account_fnp_id", 26, "fnp"),
)

_logger = logging.getLogger(__name__)


def _clean(value):
    if value is None:
        return None
    value = str(value).strip()
    return value or None


def load_rows(path: Path = XLSX_PATH) -> list[tuple]:
    wb = openpyxl.load_workbook(path, read_only=True, data_only=True)
    ws = wb[wb.sheetnames[0]]
    rows = []
    for i, row in enumerate(ws.iter_rows(values_only=True), start=1):
        if i < FIRST_DATA_ROW or not row or not _clean(row[0]):
            continue
        rows.append(row)
    if not rows:
        raise ValueError("Aucune ligne à importer")
    return rows


def _partner_vals(env, row, countries, currencies):
    name = _clean(row[0])
    ref = _clean(row[1])
    mobile = _clean(row[2])
    phone = _clean(row[3])
    email = _clean(row[4])
    currency_code = _clean(row[7])
    vat = _clean(row[8])
    patente = _clean(row[9])
    ice = _clean(row[10])
    city = _clean(row[13])
    street = _clean(row[14])
    country_name = (_clean(row[15]) or "").upper()

    phones = [p for p in (phone, mobile) if p]
    if email:
        email = re.sub(r"\s*[/;]\s*", ", ", email)
    vals = {
        "name": name,
        "ref": ref,
        "is_company": True,
        "company_type": "company",
        "supplier_rank": 1,
        "phone": " / ".join(phones) or False,
        "email": email or False,
        "vat": vat or False,
        "company_registry": ice or False,
        "rmg_l10n_ma_professional_tax_id": patente or False,
        "city": city or False,
        "street": street or False,
    }
    country = countries.get(COUNTRY_CODES.get(country_name, ""))
    if country:
        vals["country_id"] = country.id
    return vals, currencies.get(currency_code)


def run(env, commit=False, limit=None):
    company = env.company
    Account = env["account.account"].with_company(company).with_context(
        tracking_disable=True, mail_create_nolog=True, mail_notrack=True
    )
    Partner = env["res.partner"].with_context(
        tracking_disable=True, mail_create_nolog=True, mail_notrack=True
    )

    rows = load_rows()
    if limit:
        rows = rows[:limit]

    countries = {
        c.code: c for c in env["res.country"].search([("code", "in", list(COUNTRY_CODES.values()))])
    }
    currencies = {c.name: c for c in env["res.currency"].with_context(active_test=False).search([])}

    # ------------------------------------------------------------------
    # 1. Comptes comptables
    # ------------------------------------------------------------------
    base_codes = {}
    for _field, start, _label in ACCOUNT_BLOCKS:
        base_codes[start] = str(rows[0][start + 2]).strip()
    base_accounts = {}
    for start, code in base_codes.items():
        acc = Account.search([("code", "=", code), ("company_ids", "in", company.id)], limit=1)
        if not acc:
            raise ValueError(f"Compte de base {code} introuvable dans {company.name}")
        base_accounts[start] = acc

    existing = {}
    for acc in Account.with_context(active_test=False).search(
        [("company_ids", "in", company.id), ("genex_agency_code", "!=", 0)]
    ):
        if acc.genex_full_code:
            existing[acc.genex_full_code] = acc

    def full_code(row, start):
        agency, curr, code, cif, seq = (row[start + i] for i in range(5))
        return ".".join(str(int(x) if x is not None else 0) for x in (agency, curr, code, cif, seq))

    to_create = []
    seen = set()
    for row in rows:
        for _field, start, _label in ACCOUNT_BLOCKS:
            fc = full_code(row, start)
            if fc in existing or fc in seen:
                continue
            seen.add(fc)
            base = base_accounts[start]
            to_create.append(
                {
                    "name": base.with_context(lang=None).name,
                    "code": str(int(row[start + 2])),
                    "account_type": base.account_type,
                    "reconcile": base.reconcile,
                    "company_ids": [(6, 0, [company.id])],
                    "genex_agency_code": int(row[start]),
                    "genex_currency_code": int(row[start + 1]),
                    "genex_cif": int(row[start + 3] or 0),
                    "genex_sequence": int(row[start + 4] or 0),
                }
            )

    created_accounts = 0
    for i in range(0, len(to_create), BATCH_SIZE):
        batch = Account.create(to_create[i : i + BATCH_SIZE])
        for acc in batch:
            existing[acc.genex_full_code] = acc
        created_accounts += len(batch)
        if commit:
            env.cr.commit()
        env.invalidate_all()

    # ------------------------------------------------------------------
    # 2. Fournisseurs
    # ------------------------------------------------------------------
    created_partners = updated_partners = 0
    errors = []
    invalid_vat = []
    for n, row in enumerate(rows, start=1):
        ref = _clean(row[1])
        try:
            with env.cr.savepoint():
                vals, currency = _partner_vals(env, row, countries, currencies)
                partner = Partner.search([("ref", "=", ref)], limit=1) if ref else Partner
                if not partner:
                    partner = Partner.search(
                        [("name", "=ilike", vals["name"]), ("supplier_rank", ">", 0)], limit=1
                    )
                if vals.get("vat") and not re.fullmatch(r"\d{7,8}", vals["vat"].replace(" ", "")):
                    # IF invalide selon la validation RMG : on le conserve en note
                    invalid_vat.append(f"{ref} {row[0]}: IF {vals['vat']}")
                    vals["comment"] = f"IF (référentiel, non conforme) : {vals['vat']}"
                    vals["vat"] = False
                if partner:
                    partner.write(vals)
                    updated_partners += 1
                else:
                    partner = Partner.create(vals)
                    created_partners += 1
                scoped = partner.with_company(company)
                prop_vals = {
                    field: existing[full_code(row, start)].id
                    for field, start, _label in ACCOUNT_BLOCKS
                }
                if currency:
                    prop_vals["property_purchase_currency_id"] = currency.id
                scoped.write(prop_vals)
        except Exception as exc:  # noqa: BLE001
            errors.append(f"{ref} {row[0]}: {exc}")
        if commit and n % BATCH_SIZE == 0:
            env.cr.commit()
    if commit:
        env.cr.commit()

    return {
        "rows": len(rows),
        "accounts_created": created_accounts,
        "accounts_reused": len(existing) - created_accounts,
        "partners_created": created_partners,
        "partners_updated": updated_partners,
        "invalid_vat": invalid_vat,
        "errors": errors,
    }


if __name__ == "__main__":
    print("Ce script doit être exécuté via odoo-bin shell.", file=sys.stderr)
    sys.exit(1)
